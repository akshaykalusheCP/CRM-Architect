"""Text extraction from business documents (SOPs, invoices, quotations, forms)."""

import io
from pathlib import Path

DOCUMENT_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".json"}


def extract_text(filename: str, data: bytes) -> str:
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        pages = [(page.extract_text() or "").strip() for page in reader.pages]
        text = "\n\n".join(f"[page {i + 1}]\n{p}" for i, p in enumerate(pages) if p)
        if not text.strip():
            raise ValueError("PDF has no extractable text (it may be a scanned image; OCR is not enabled)")
        return text
    if ext == ".docx":
        import docx

        document = docx.Document(io.BytesIO(data))
        parts = [p.text for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        return "\n".join(parts)
    if ext in {".txt", ".md", ".json"}:
        for encoding in ("utf-8-sig", "latin-1"):
            try:
                return data.decode(encoding)
            except UnicodeDecodeError:
                continue
    raise ValueError(f"Unsupported document type: {ext}")
